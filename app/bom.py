from pathlib import Path

import openpyxl

from .models import BomRow
from .tabular import read_csv_rows

REFERENCE_COLUMN = "Nº de pieza"
MATERIAL_COLUMN = "Material"
CATEGORY_COLUMN = "Categoría"
TARGET_CATEGORY = "LASER"


def parse_bom(path: str | Path, sheet_name: str = "BOM") -> list[BomRow]:
    """Parse a CM4 BOM (`.xlsx`, or `.csv` with the same column names), returning one row
    per LASER-category part that has a specified material.

    Rows with no material (assemblies, e.g. "PISADERA RENFE HC EN EMBALAJE") are excluded:
    they have nothing to verify against a certificate. Rows outside the LASER category
    (MECANIZADO, ZINCADO, MONTAJE, ...) are also excluded - out of scope per current
    verification focus. A BOM with no `Categoría` column at all is not filtered by
    category (keeps this usable against simpler/synthetic BOMs).
    """
    path = Path(path)
    if path.suffix.lower() == ".csv":
        rows = read_csv_rows(path.read_bytes())
        source = "BOM CSV"
    else:
        wb = openpyxl.load_workbook(path, data_only=True)
        # A BOM exported on its own often has a single, differently named sheet.
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        source = f"BOM sheet '{ws.title}'"
    return _rows_to_bom(rows, source)


def _rows_to_bom(rows: list, source: str) -> list[BomRow]:
    if not rows:
        raise ValueError(f"{source} is empty")
    header = [str(name).strip() if name is not None else None for name in rows[0]]
    col_index = {name: i for i, name in enumerate(header) if name}

    if REFERENCE_COLUMN not in col_index or MATERIAL_COLUMN not in col_index:
        raise ValueError(
            f"{source} is missing expected columns "
            f"'{REFERENCE_COLUMN}' / '{MATERIAL_COLUMN}'"
        )

    ref_idx = col_index[REFERENCE_COLUMN]
    mat_idx = col_index[MATERIAL_COLUMN]
    cat_idx = col_index.get(CATEGORY_COLUMN)

    result: list[BomRow] = []
    for row in rows[1:]:
        if row is None or ref_idx >= len(row) or mat_idx >= len(row):
            continue
        reference = row[ref_idx]
        material = row[mat_idx]
        if reference is None or (isinstance(reference, str) and not reference.strip()):
            continue
        if material is None or (isinstance(material, str) and not material.strip()):
            continue
        if cat_idx is not None:
            category = row[cat_idx] if cat_idx < len(row) else None
            if not category or str(category).strip().upper() != TARGET_CATEGORY:
                continue
        result.append(BomRow(reference=str(reference).strip(), material=str(material).strip()))

    return result
