import io
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from .i18n import DEFAULT_LANG, translate
from .models import OrphanCorrespondenceEntry, Status, VerificationRecord

# Stable field keys, in column order. Headers are rendered from the catalog
# ("export.h.<key>") in the export's language; app/history.py maps them back.
EXPORT_FIELDS = [
    "reference",
    "specified_norma",
    "specified_material",
    "extracted_material",
    "certificate",
    "certificate_type",
    "certificate_type_warning",
    "committed_material",
    "status",
    "reason",
    "suggested_action",
    "human_confirmed",
]
ORPHAN_FIELDS = ["reference", "lote"]
SUMMARY_FIELDS = [
    "verified_at",
    "exported_at",
    "run_id",
    "bom",
    "certificates",
    "suppliers",
    "correspondence",
    "total",
    *(f"count.{status.value}" for status in Status),
    "human_confirmed",
    "orphans",
]

EXPORT_HEADERS = [translate(DEFAULT_LANG, f"export.h.{key}") for key in EXPORT_FIELDS]
ORPHANS_HEADERS = [translate(DEFAULT_LANG, f"export.orphans.h.{key}") for key in ORPHAN_FIELDS]


@dataclass
class RunMetadata:
    run_id: str
    created_at: datetime
    bom_filename: str | None = None
    certificate_count: int = 0
    suppliers: list[str] = field(default_factory=list)
    correspondence_filename: str | None = None


def export_filename(metadata: RunMetadata, lang: str | None = DEFAULT_LANG) -> str:
    prefix = translate(lang, "export.filename_prefix")
    return f"{prefix}_{metadata.created_at:%Y-%m-%d_%H%M}_{metadata.run_id[:8]}.xlsx"


def _record_row(record: VerificationRecord, lang: str) -> list:
    action = (
        translate(lang, record.action_key, **record.action_params)
        if record.action_key
        else record.suggested_action
    )
    warning = (
        translate(lang, record.warning_key, **record.warning_params)
        if record.warning_key
        else record.certificate_type_warning
    )
    return [
        record.reference,
        record.specified_norma,
        record.specified_material,
        record.extracted_material,
        record.certificate_filename,
        record.certificate_type,
        warning,
        record.committed_material,
        record.status.value,  # raw code on purpose: language-independent for app/history.py
        record.reason.value if record.reason else None,
        action,
        record.human_confirmed,
    ]


def build_export_workbook(
    records: list[VerificationRecord],
    orphan_correspondence_entries: list[OrphanCorrespondenceEntry] | None = None,
    *,
    lang: str | None = DEFAULT_LANG,
    metadata: RunMetadata | None = None,
    exported_at: datetime | None = None,
) -> bytes:
    """Build the export workbook: one row per verified part, with every decision, the
    certificate involved, and the materials asked for vs. produced. A second sheet lists
    correspondence-document rows that matched no BOM reference at all (see
    matching.find_orphan_correspondence_entries), when there are any. With `metadata`, a
    summary sheet records when the run happened, from which inputs, and its per-status counts,
    so the file is self-describing in the run history (app/history.py).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = translate(lang, "export.sheet.results")
    ws.append([translate(lang, f"export.h.{key}") for key in EXPORT_FIELDS])
    for record in records:
        ws.append(_record_row(record, lang))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    if orphan_correspondence_entries:
        orphans_ws = wb.create_sheet(translate(lang, "export.sheet.orphans"))
        orphans_ws.append([translate(lang, f"export.orphans.h.{key}") for key in ORPHAN_FIELDS])
        for entry in orphan_correspondence_entries:
            orphans_ws.append([entry.reference, entry.lote])

    if metadata is not None:
        counts = {status.value: 0 for status in Status}
        for record in records:
            counts[record.status.value] += 1
        values = {
            "verified_at": metadata.created_at.replace(microsecond=0),
            "exported_at": (exported_at or datetime.now()).replace(microsecond=0),
            "run_id": metadata.run_id,
            "bom": metadata.bom_filename,
            "certificates": metadata.certificate_count,
            "suppliers": ", ".join(metadata.suppliers),
            "correspondence": metadata.correspondence_filename,
            "total": len(records),
            **{f"count.{k}": v for k, v in counts.items()},
            "human_confirmed": sum(1 for r in records if r.human_confirmed),
            "orphans": len(orphan_correspondence_entries or []),
        }
        summary_ws = wb.create_sheet(translate(lang, "export.sheet.summary"))
        for key in SUMMARY_FIELDS:
            summary_ws.append([translate(lang, f"summary.{key}"), values[key]])
        for row in summary_ws.iter_rows(min_col=1, max_col=1):
            row[0].font = Font(bold=True)
        summary_ws.column_dimensions["A"].width = 30
        summary_ws.column_dimensions["B"].width = 40

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def save_export(content: bytes, directory: Path, filename: str) -> Path:
    """Write an export atomically (temp file + rename), so the history folder never holds a
    half-written workbook even if the app stops mid-save."""
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / filename
    tmp = directory / f".{filename}.tmp"
    tmp.write_bytes(content)
    tmp.replace(target)
    return target
