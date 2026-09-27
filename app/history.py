"""Run history: reads back the Excel exports saved in the exports folder (config.EXPORTS_DIR).

The folder is the history - there is no separate store. Exports may be written in Catalan or
Spanish, or predate the summary sheet, so every sheet and column is recognised by its header
text in *any* catalog language (see i18n.labels_for) rather than by sheet name or position.
"""

import io
import zipfile
from datetime import date, datetime
from pathlib import Path

import openpyxl

from .export import EXPORT_FIELDS, ORPHAN_FIELDS, SUMMARY_FIELDS
from .i18n import labels_for
from .models import Status


def _label_map(prefix: str, keys: list[str]) -> dict[str, str]:
    return {label.strip().lower(): key for key in keys for label in labels_for(f"{prefix}{key}")}


_RECORD_HEADERS = _label_map("export.h.", EXPORT_FIELDS)
_ORPHAN_HEADERS = _label_map("export.orphans.h.", ORPHAN_FIELDS)
_SUMMARY_LABELS = _label_map("summary.", SUMMARY_FIELDS)


class NotAnExport(ValueError):
    pass


def _json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _header_keys(row: tuple, mapping: dict[str, str]) -> list[str | None]:
    return [mapping.get(str(cell).strip().lower()) if cell is not None else None for cell in row]


def _read_table(ws, mapping: dict[str, str], required: set[str]) -> list[dict] | None:
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return None
    keys = _header_keys(header, mapping)
    if not required <= set(keys):
        return None
    out = []
    for row in rows:
        if row is None or all(cell is None for cell in row):
            continue
        out.append({k: _json_value(v) for k, v in zip(keys, row) if k is not None})
    return out


def _read_summary(ws) -> dict | None:
    values = {}
    for row in ws.iter_rows(values_only=True, max_col=2):
        if not row or row[0] is None:
            continue
        key = _SUMMARY_LABELS.get(str(row[0]).strip().lower())
        if key:
            values[key] = row[1] if len(row) > 1 else None
    # A sheet only counts as the summary if it at least says when the run happened.
    return values if "verified_at" in values else None


def _as_datetime(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return None
    return None


def parse_export(content: bytes, filename: str, modified_at: datetime) -> dict:
    """One export file as a history run. Raises NotAnExport when no sheet looks like a
    verification results sheet."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise NotAnExport("unreadable") from exc

    records = orphans = summary = None
    for ws in wb.worksheets:
        if records is None:
            records = _read_table(ws, _RECORD_HEADERS, {"reference", "status"})
            if records is not None:
                continue
        if orphans is None:
            orphans = _read_table(ws, _ORPHAN_HEADERS, {"reference", "lote"})
            if orphans is not None:
                continue
        if summary is None:
            summary = _read_summary(ws)
    wb.close()
    if records is None:
        raise NotAnExport("not_export")

    counts = {status.value: 0 for status in Status}
    for record in records:
        if record.get("status") in counts:
            counts[record["status"]] += 1

    summary = summary or {}
    verified_at = _as_datetime(summary.get("verified_at"))
    exported_at = _as_datetime(summary.get("exported_at"))
    run_date = verified_at or modified_at
    certificates = summary.get("certificates")
    return {
        "file": filename,
        "date": run_date.isoformat(timespec="minutes"),
        "date_is_approximate": verified_at is None,
        "exported_at": exported_at.isoformat(timespec="minutes") if exported_at else None,
        "run_id": summary.get("run_id"),
        "bom": summary.get("bom"),
        "certificates": certificates if isinstance(certificates, int) else None,
        "suppliers": summary.get("suppliers"),
        "correspondence": summary.get("correspondence"),
        "total": len(records),
        "summary": counts,
        "human_confirmed": sum(1 for r in records if r.get("human_confirmed") is True),
        "records": records,
        "orphans": orphans or [],
    }


def list_history(directory: Path) -> dict:
    """Every export under `directory`, newest first, plus files that were skipped and why."""
    runs, ignored = [], []
    if directory.is_dir():
        for path in sorted(directory.rglob("*.xlsx")):
            relative = str(path.relative_to(directory))
            if path.name.startswith(("~$", ".")):
                continue  # Excel lock files and our own in-progress temp files
            modified_at = datetime.fromtimestamp(path.stat().st_mtime)
            try:
                runs.append(parse_export(path.read_bytes(), relative, modified_at))
            except NotAnExport as exc:
                ignored.append({"file": relative, "reason": str(exc)})
    runs.sort(key=lambda run: run["date"], reverse=True)
    return {"folder": str(directory), "runs": runs, "ignored": ignored}
