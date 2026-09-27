import os
from datetime import datetime
from pathlib import Path

import openpyxl

from app.export import build_export_workbook, save_export
from app.history import list_history
from test_export import _metadata, _two_records

def test_reads_spanish_and_catalan_exports_newest_first(tmp_path):
    older = _metadata()
    newer = _metadata()
    newer.created_at = datetime(2026, 9, 28, 9, 0)
    newer.run_id = "ffff0000"
    save_export(build_export_workbook(_two_records(), lang="ca", metadata=older), tmp_path, "a.xlsx")
    save_export(build_export_workbook(_two_records(), lang="es", metadata=newer), tmp_path, "b.xlsx")

    result = list_history(tmp_path)

    assert [run["file"] for run in result["runs"]] == ["b.xlsx", "a.xlsx"]
    run = result["runs"][1]
    assert run["date"] == "2026-09-27T14:30"
    assert run["date_is_approximate"] is False
    assert run["bom"] == "M009_BOM.xlsx"
    assert run["certificates"] == 11
    assert run["summary"] == {"ok": 1, "needs_review": 0, "mismatch": 0, "missing_certificate": 1}
    assert run["human_confirmed"] == 1
    assert run["records"][0]["reference"] == "A-1"
    assert result["runs"][0]["records"][0]["suggested_action"].startswith("Guardar como")


def _legacy_export(path: Path) -> None:
    """An export as written before summary sheets and i18n existed: Catalan headers, a results
    sheet and an orphans sheet, nothing else."""
    from app.export import EXPORT_HEADERS, ORPHANS_HEADERS

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Verificació (LASER)"
    ws.append(EXPORT_HEADERS)
    ws.append(["A-1", "EN 10025-2:2004", "EN 10025-2:2004 S235JR", "EN10025-2 S235JR", "S/1/C1.pdf", "2.2",
               "Certificat de tipus 2.2 (no 3.1)", None, "needs_review", "certificate_type", "Revisar", False])
    ws.append(["B-2", None, "EN AW 5754", "EN AW-5754", "S/2/C2.pdf", None, None, None, "ok", None, "Guardar", False])
    ws.append(["C-3", None, "EN AW 6063", None, None, None, None, None, "missing_certificate", None, "Reclamar", False])
    orphans = wb.create_sheet("Correspondència sense BOM")
    orphans.append(ORPHANS_HEADERS)
    orphans.append(["Z-9", "L123"])
    wb.save(path)


def test_legacy_export_without_summary_uses_file_date(tmp_path):
    target = tmp_path / "old" / "verificacio_4bc7a973.xlsx"
    target.parent.mkdir()
    _legacy_export(target)
    stamp = datetime(2026, 9, 20, 10, 15).timestamp()
    os.utime(target, (stamp, stamp))

    run = list_history(tmp_path)["runs"][0]

    assert run["file"] == "old/verificacio_4bc7a973.xlsx"
    assert run["date"] == "2026-09-20T10:15"
    assert run["date_is_approximate"] is True
    assert run["total"] == 3
    assert run["summary"] == {"ok": 1, "needs_review": 1, "mismatch": 0, "missing_certificate": 1}
    assert run["orphans"] == [{"reference": "Z-9", "lote": "L123"}]


def test_non_exports_are_ignored_with_a_reason_and_lock_files_skipped(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["Nº de pieza", "Material"])
    wb.save(tmp_path / "some_bom.xlsx")
    (tmp_path / "broken.xlsx").write_bytes(b"not a workbook")
    (tmp_path / "~$open.xlsx").write_bytes(b"lock")
    save_export(build_export_workbook(_two_records(), metadata=_metadata()), tmp_path, "run.xlsx")

    result = list_history(tmp_path)

    assert [run["file"] for run in result["runs"]] == ["run.xlsx"]
    assert {i["file"]: i["reason"] for i in result["ignored"]} == {
        "broken.xlsx": "unreadable",
        "some_bom.xlsx": "not_export",
    }


def test_missing_folder_is_an_empty_history(tmp_path):
    result = list_history(tmp_path / "nope")
    assert result["runs"] == [] and result["ignored"] == []
