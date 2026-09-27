from app.bom import parse_bom


def test_parses_reference_and_material(examples_dir):
    rows = parse_bom(examples_dir / "EBRO" / "M009_BOM.xlsx")
    by_ref = {r.reference: r.material for r in rows}
    # I018-L106 is category LASER in the sample BOM.
    assert by_ref["I018-L106"] == "EN 10025-2:2004 S235JR"


def test_excludes_rows_with_no_material(examples_dir):
    rows = parse_bom(examples_dir / "EBRO" / "M009_BOM.xlsx")
    refs = {r.reference for r in rows}
    # M009-0001 and M009-0000 are assemblies with an empty Material column in the sample BOM
    assert "M009-0001" not in refs
    assert "M009-0000" not in refs


def test_excludes_rows_outside_laser_category(examples_dir):
    rows = parse_bom(examples_dir / "EBRO" / "M009_BOM.xlsx")
    refs = {r.reference for r in rows}
    # I018-M108 has a material but is category MECANIZADO, not LASER - out of scope.
    assert "I018-M108" not in refs


def _bom_as_csv(examples_dir, delimiter: str, encoding: str) -> bytes:
    import csv
    import io

    import openpyxl

    ws = openpyxl.load_workbook(examples_dir / "EBRO" / "M009_BOM.xlsx", data_only=True)["BOM"]
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=delimiter)
    for row in ws.iter_rows(values_only=True):
        writer.writerow(["" if v is None else v for v in row])
    return buffer.getvalue().encode(encoding)


def test_csv_bom_parses_like_the_xlsx(examples_dir, tmp_path):
    expected = parse_bom(examples_dir / "EBRO" / "M009_BOM.xlsx")
    path = tmp_path / "bom.csv"
    path.write_bytes(_bom_as_csv(examples_dir, ",", "utf-8-sig"))
    assert parse_bom(path) == expected


def test_semicolon_cp1252_csv_bom_as_excel_es_writes_it(examples_dir, tmp_path):
    # Spanish/Catalan-locale Excel saves CSV with ";" and Windows-1252 ("Nº", "Categoría").
    expected = parse_bom(examples_dir / "EBRO" / "M009_BOM.xlsx")
    path = tmp_path / "bom.csv"
    path.write_bytes(_bom_as_csv(examples_dir, ";", "cp1252"))
    assert parse_bom(path) == expected


def test_xlsx_without_bom_sheet_uses_first_sheet(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.append(["Nº de pieza", "Material"])
    ws.append(["X-1", "S235JR"])
    path = tmp_path / "bom.xlsx"
    wb.save(path)
    assert [(r.reference, r.material) for r in parse_bom(path)] == [("X-1", "S235JR")]


def test_csv_missing_columns_is_a_clear_error(tmp_path):
    import pytest

    path = tmp_path / "bom.csv"
    path.write_text("Ref;Mat\nA;B\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Nº de pieza"):
        parse_bom(path)
