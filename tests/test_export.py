import io

import openpyxl

from app.export import EXPORT_HEADERS, ORPHANS_HEADERS, build_export_workbook
from app.models import OrphanCorrespondenceEntry, Status, VerificationRecord


def test_export_workbook_has_expected_headers_and_rows():
    records = [
        VerificationRecord(
            reference="I018-L106",
            specified_material="EN 10025-2:2004 S235JR",
            specified_norma="EN 10025-2:2004",
            specified_grade="S235JR",
            extracted_material="EN 10025-2 S235JR",
            certificate_filename="EBRO/11703780001/ZA34852.pdf",
            status=Status.OK,
            certificate_type="2.2",
            certificate_type_warning="Certificat de tipus 2.2 (no 3.1): revisar traçabilitat.",
        ),
        VerificationRecord(
            reference="RCM4-L082",
            specified_material="EN 10088-1:1995 X5CrNi18-10",
            extracted_material=None,
            certificate_filename=None,
            status=Status.MISSING_CERTIFICATE,
        ),
    ]

    workbook_bytes = build_export_workbook(records)
    wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes))
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == EXPORT_HEADERS

    by_ref = {row[0]: row for row in rows[1:]}
    ok_row = by_ref["I018-L106"]
    assert ok_row[1] == "EN 10025-2:2004"  # Norma especificada
    assert ok_row[2] == "EN 10025-2:2004 S235JR"  # Material especificat
    assert ok_row[3] == "EN 10025-2 S235JR"  # Material de fabricació
    assert ok_row[5] == "2.2"  # Tipus de certificat
    assert ok_row[8] == "ok"  # Resultat

    missing_row = by_ref["RCM4-L082"]
    assert missing_row[4] is None  # Certificat associat
    assert missing_row[8] == "missing_certificate"

    assert wb.sheetnames == ["Verificació (LASER)"]  # no orphans sheet when there are none


def test_export_workbook_adds_orphans_sheet_when_present():
    records = [
        VerificationRecord(
            reference="I018-L106",
            specified_material="EN 10025-2:2004 S235JR",
            extracted_material="EN 10025-2 S235JR",
            certificate_filename="EBRO/x/ZA1.pdf",
            status=Status.OK,
        )
    ]
    orphans = [OrphanCorrespondenceEntry(reference="M009-L114", lote="ZA33404")]

    workbook_bytes = build_export_workbook(records, orphans)
    wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes))
    assert "Correspondència sense BOM" in wb.sheetnames

    orphans_ws = wb["Correspondència sense BOM"]
    rows = list(orphans_ws.iter_rows(values_only=True))
    assert list(rows[0]) == ORPHANS_HEADERS
    assert rows[1] == ("M009-L114", "ZA33404")


def _metadata():
    from datetime import datetime

    from app.export import RunMetadata

    return RunMetadata(
        run_id="ab12cd34-0000-0000-0000-000000000000",
        created_at=datetime(2026, 9, 27, 14, 30, 5),
        bom_filename="M009_BOM.xlsx",
        certificate_count=11,
        suppliers=["EBRO"],
        correspondence_filename="correspondencia_1170378.xlsx",
    )


def _two_records():
    return [
        VerificationRecord(
            reference="A-1", specified_material="S235JR", extracted_material="S235JR",
            certificate_filename="EBRO/x/ZA1.pdf", status=Status.OK,
            action_key="action.save_as", action_params={"filename": "A-1_ZA1_S235JR.pdf"},
            suggested_action="Guardar com A-1_ZA1_S235JR.pdf",
        ),
        VerificationRecord(
            reference="B-2", specified_material="S355JR", extracted_material=None,
            certificate_filename=None, status=Status.MISSING_CERTIFICATE, human_confirmed=True,
            action_key="action.chase_supplier", action_params={"reference": "B-2"},
        ),
    ]


def test_export_in_spanish_with_summary_sheet():
    from app.export import export_filename

    wb = openpyxl.load_workbook(io.BytesIO(build_export_workbook(_two_records(), lang="es", metadata=_metadata())))
    assert wb.sheetnames == ["Verificación (LASER)", "Resumen"]
    rows = list(wb["Verificación (LASER)"].iter_rows(values_only=True))
    assert rows[0][0] == "Referencia" and rows[0][8] == "Resultado"
    assert rows[1][10] == "Guardar como A-1_ZA1_S235JR.pdf"
    assert rows[2][8] == "missing_certificate"  # status codes stay language-independent

    summary = {row[0]: row[1] for row in wb["Resumen"].iter_rows(values_only=True)}
    assert summary["Archivo BOM"] == "M009_BOM.xlsx"
    assert summary["Total de piezas"] == 2
    assert summary["Certificado no encontrado"] == 1
    assert summary["Confirmados manualmente"] == 1
    assert export_filename(_metadata(), "es") == "verificacion_2026-09-27_1430_ab12cd34.xlsx"
    assert export_filename(_metadata(), "ca") == "verificacio_2026-09-27_1430_ab12cd34.xlsx"
